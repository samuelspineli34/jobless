"""Gera currículos (DOCX + PDF) em vários idiomas a partir do perfil.
Corrige problemas de quebra de página, caracteres especiais e fontes."""
import re
from xml.sax.saxutils import escape

from ..services.skills import display, extract_skills

LABELS = {
    "pt": dict(summary="Resumo", experience="Experiência Profissional", education="Formação", skills="Competências", languages="Idiomas", present="Atual", links="Links"),
    "en": dict(summary="Summary", experience="Professional Experience", education="Education", skills="Skills", languages="Languages", present="Present", links="Links"),
    "es": dict(summary="Resumen", experience="Experiencia Profesional", education="Formación", skills="Competencias", languages="Idiomas", present="Actual", links="Enlaces"),
    "fr": dict(summary="Résumé", experience="Expérience Professionnelle", education="Formation", skills="Compétences", languages="Langues", present="Présent", links="Liens"),
    "de": dict(summary="Zusammenfassung", experience="Berufserfahrung", education="Ausbildung", skills="Fähigkeiten", languages="Sprachen", present="Heute", links="Links"),
    "it": dict(summary="Profilo", experience="Esperienza Professionale", education="Istruzione", skills="Competenze", languages="Lingue", present="Presente", links="Link"),
}
LANG_NAMES = {
    "pt": {"pt": "Português", "en": "Inglês", "es": "Espanhol", "fr": "Francês", "de": "Alemão", "it": "Italiano"},
    "en": {"pt": "Portuguese", "en": "English", "es": "Spanish", "fr": "French", "de": "German", "it": "Italian"}
}
SUPPORTED = list(LABELS)


def _safe_text(val: str | None) -> str:
    if not val:
        return ""
    return escape(str(val).strip())


def _tr(obj: dict, field: str, lang: str, base_lang: str, warnings: list[str]) -> str:
    tr = (obj.get("i18n") or {}).get(lang, {}).get(field)
    if tr:
        return tr
    val = obj.get(field) or ""
    if val and lang != base_lang:
        warnings.append(field)
    return val


def _fmt(d: str | None, present: str) -> str:
    if not d:
        return present
    m = re.match(r"(\d{4})(?:-(\d{2}))?", d)
    return f"{m.group(2)}/{m.group(1)}" if m and m.group(2) else (m.group(1) if m else d)


def build_content(profile: dict, experiences: list[dict], education: list[dict], skills: list[str],
                  lang: str, job_skills: list[str] | None = None) -> dict:
    if lang not in LABELS:
        lang = "en"
    L = LABELS[lang]
    base = (profile.get("preferences") or {}).get("content_language", "pt")
    warns: list[str] = []
    jobset = {s.lower() for s in (job_skills or [])}
    skills_sorted = sorted(skills, key=lambda s: (s.lower() not in jobset, s.lower()))
    exps = sorted(experiences, key=lambda e: e.get("start") or "", reverse=True)
    exp_out = []
    for e in exps:
        desc = _tr(e, "description", lang, base, warns)
        bullets = [re.sub(r"^[\-•*·]\s*", "", l).strip() for l in desc.splitlines() if l.strip()]
        exp_out.append({
            "title": e.get("title", ""),
            "company": e.get("company", ""),
            "location": e.get("location", ""),
            "period": f"{_fmt(e.get('start'), '')} – {_fmt(e.get('end'), L['present'])}",
            "bullets": bullets,
            "techs": [display(t) for t in (e.get("technologies") or [])]
        })

    langs = [f"{LANG_NAMES.get(lang, LANG_NAMES['en']).get(x['code'], x['code'])}" + (f" ({x['level']})" if x.get("level") else "")
             for x in profile.get("languages", [])]
    contact = [x for x in (profile.get("email"), profile.get("phone"), profile.get("location")) if x]
    links = [v for v in (profile.get("links") or {}).values() if v]

    content = {
        "labels": L, "name": profile.get("full_name", ""), "headline": _tr(profile, "headline", lang, base, warns),
        "contact": contact, "links": links, "summary": _tr(profile, "summary", lang, base, warns),
        "experiences": exp_out, "skills": [display(s) for s in skills_sorted[:30]], "languages": langs,
        "education": [{"institution": e.get("institution", ""), "degree": _tr(e, "degree", lang, base, warns),
                       "period": " – ".join(x for x in (_fmt(e.get("start"), ""), _fmt(e.get("end"), "")) if x)} for e in education],
    }
    content["warnings"] = sorted(set(warns))
    return content


def flat_text(c: dict) -> str:
    parts = [c["name"], c["headline"], c["summary"], " ".join(c["skills"])]
    for e in c["experiences"]:
        parts += [e["title"], e["company"], " ".join(e["bullets"]), " ".join(e["techs"])]
    return "\n".join(p for p in parts if p)


def find_unverified(c: dict, profile_skills: list[str], profile_texts: list[str]) -> list[str]:
    known = {s.lower() for s in profile_skills} | set(extract_skills(" ".join(profile_texts)))
    return sorted(set(extract_skills(flat_text(c))) - known)


def write_docx(c: dict, path: str) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)

    doc.add_heading(c["name"], level=0)
    if c["headline"]:
        p = doc.add_paragraph(c["headline"])
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("  |  ".join(c["contact"] + c["links"]))
    L = c["labels"]

    if c["summary"]:
        h = doc.add_heading(L["summary"], level=1)
        h.paragraph_format.keep_with_next = True
        doc.add_paragraph(c["summary"])

    if c["experiences"]:
        h = doc.add_heading(L["experience"], level=1)
        h.paragraph_format.keep_with_next = True
        for e in c["experiences"]:
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{e['title']} — {e['company']}")
            r.bold = True
            p.add_run(f"   ({e['period']})" + (f" · {e['location']}" if e["location"] else ""))
            for b in e["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
            if e["techs"]:
                doc.add_paragraph("Stack: " + ", ".join(e["techs"]))

    if c["education"]:
        h = doc.add_heading(L["education"], level=1)
        h.paragraph_format.keep_with_next = True
        for e in c["education"]:
            doc.add_paragraph(f"{e['institution']} — {e['degree']}  {e['period']}".strip())

    if c["skills"]:
        h = doc.add_heading(L["skills"], level=1)
        h.paragraph_format.keep_with_next = True
        doc.add_paragraph(", ".join(c["skills"]))

    if c["languages"]:
        h = doc.add_heading(L["languages"], level=1)
        h.paragraph_format.keep_with_next = True
        doc.add_paragraph(", ".join(c["languages"]))

    doc.save(path)


def write_pdf(c: dict, path: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9, leading=12.5, spaceAfter=2)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=7, spaceAfter=3,
                        textColor=colors.HexColor("#1b365d"), keepWithNext=True)
    name = ParagraphStyle("n", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=2)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    bullet = ParagraphStyle("bl", parent=body, leftIndent=10, bulletIndent=2, spaceAfter=1)

    story = [Paragraph(_safe_text(c["name"]), name)]
    if c["headline"]:
        story.append(Paragraph(_safe_text(c["headline"]), body))

    contact_line = "  |  ".join(_safe_text(x) for x in (c["contact"] + c["links"]))
    story.append(Paragraph(contact_line, small))
    story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#1b365d"), spaceBefore=4, spaceAfter=6))
    L = c["labels"]

    def section(title):
        story.append(Paragraph(_safe_text(title).upper(), h1))

    if c["summary"]:
        section(L["summary"])
        story.append(Paragraph(_safe_text(c["summary"]), body))

    if c["experiences"]:
        section(L["experience"])
        for e in c["experiences"]:
            title_line = f"<b>{_safe_text(e['title'])}</b> — {_safe_text(e['company'])} <font color='#555555'>({_safe_text(e['period'])})</font>"
            story.append(Paragraph(title_line, body))
            for b in e["bullets"]:
                story.append(Paragraph(_safe_text(b), bullet, bulletText="•"))
            if e["techs"]:
                story.append(Paragraph(f"<i>Stack:</i> {', '.join(_safe_text(t) for t in e['techs'])}", small))
            story.append(Spacer(1, 3))

    if c["education"]:
        section(L["education"])
        for e in c["education"]:
            story.append(Paragraph(f"{_safe_text(e['institution'])} — {_safe_text(e['degree'])}  {_safe_text(e['period'])}".strip(), body))

    if c["skills"]:
        section(L["skills"])
        story.append(Paragraph(", ".join(_safe_text(s) for s in c["skills"]), body))

    if c["languages"]:
        section(L["languages"])
        story.append(Paragraph(", ".join(_safe_text(l) for l in c["languages"]), body))

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                            title=f"{c['name']} - CV", author=c["name"])
    doc.build(story)